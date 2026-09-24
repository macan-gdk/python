"""Xuất nội dung ra .txt / .docx."""

from __future__ import annotations

import io
from pathlib import Path

from doc_tool.readers import read_attachment
from doc_tool.types import ProgressFn


def write_docx_text(content: str, output_path: Path) -> Path:
    from docx import Document

    doc = Document()
    for line in content.splitlines():
        doc.add_paragraph(line)
    doc.save(str(output_path))
    return output_path


def write_docx_visual(
    pdf_path: Path,
    output_path: Path,
    *,
    max_pages: int | None = None,
    scale: float = 2.0,
    progress: ProgressFn | None = None,
) -> Path:
    """Mỗi trang PDF = 1 ảnh full-page trong Word."""
    import pymupdf
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    doc = Document()
    if doc.paragraphs:
        el = doc.paragraphs[0]._element
        el.getparent().remove(el)

    pdf = pymupdf.open(str(pdf_path))
    limit = min(len(pdf), max_pages) if max_pages else len(pdf)
    try:
        for i in range(limit):
            if progress:
                progress(i + 1, limit)
            page = pdf[i]
            width_in = page.rect.width / 72.0
            height_in = page.rect.height / 72.0

            section = doc.sections[0] if i == 0 else doc.add_section()
            section.page_width = Inches(width_in)
            section.page_height = Inches(height_in)
            section.left_margin = Inches(0)
            section.right_margin = Inches(0)
            section.top_margin = Inches(0)
            section.bottom_margin = Inches(0)
            if width_in > height_in:
                section.orientation = WD_ORIENT.LANDSCAPE

            pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale))
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run().add_picture(
                io.BytesIO(pix.tobytes("png")), width=Inches(width_in)
            )
    finally:
        pdf.close()

    doc.save(str(output_path))
    return output_path


def export_to_txt(
    file_path: str | Path,
    output_path: str | Path | None = None,
    *,
    use_ocr: bool = True,
    max_pages: int | None = None,
    progress: ProgressFn | None = None,
) -> Path | None:
    src = Path(file_path)
    content = read_attachment(
        src, use_ocr=use_ocr, max_pages=max_pages, progress=progress
    )
    if not content.strip():
        return None
    out = Path(output_path) if output_path else src.with_suffix(".txt")
    out.write_text(content, encoding="utf-8")
    return out


def export_to_docx(
    file_path: str | Path,
    output_path: str | Path | None = None,
    *,
    use_ocr: bool = True,
    max_pages: int | None = None,
    progress: ProgressFn | None = None,
    layout: str = "text",
) -> Path | None:
    src = Path(file_path)
    out = Path(output_path) if output_path else src.with_suffix(".docx")
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")

    layout = layout.lower().strip()
    if layout == "visual":
        if src.suffix.lower() != ".pdf":
            raise ValueError("layout='visual' chỉ hỗ trợ file PDF")
        return write_docx_visual(
            src, out, max_pages=max_pages, progress=progress
        )

    content = read_attachment(
        src, use_ocr=use_ocr, max_pages=max_pages, progress=progress
    )
    if not content.strip():
        return None
    return write_docx_text(content, out)


def export_attachment(
    file_path: str | Path,
    output_path: str | Path | None = None,
    *,
    use_ocr: bool = True,
    max_pages: int | None = None,
    progress: ProgressFn | None = None,
    layout: str = "text",
) -> Path | None:
    src = Path(file_path)
    out = Path(output_path) if output_path else src.with_suffix(".docx")
    if out.suffix.lower() == ".txt":
        return export_to_txt(
            src, out, use_ocr=use_ocr, max_pages=max_pages, progress=progress
        )
    return export_to_docx(
        src,
        out,
        use_ocr=use_ocr,
        max_pages=max_pages,
        progress=progress,
        layout=layout,
    )
