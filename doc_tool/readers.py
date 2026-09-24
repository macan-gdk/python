"""Đọc nội dung PDF / Word / Excel."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, cast

from doc_tool.tesseract_env import configure_tesseract
from doc_tool.types import ProgressFn
from doc_tool.viet_correct import fix_vietnamese_ocr

SUPPORTED_EXTS = {".pdf", ".docx", ".xlsx", ".xls"}


def _page_clip(page) -> object:
    """Cắt lề nhẹ để giảm watermark mép trang."""
    import pymupdf

    rect = page.rect
    return pymupdf.Rect(
        rect.x0 + rect.width * 0.02,
        rect.y0 + rect.height * 0.01,
        rect.x1 - rect.width * 0.08,
        rect.y1 - rect.height * 0.02,
    )


def _preprocess(img):
    import cv2
    import numpy as np
    from PIL import Image

    gray = cv2.cvtColor(np.array(img.convert("RGB")), cv2.COLOR_RGB2GRAY)
    gray = cv2.fastNlMeansDenoising(gray, None, 6, 7, 21)
    clahe = cv2.createCLAHE(clipLimit=1.5, tileGridSize=(8, 8))
    return Image.fromarray(clahe.apply(gray))


def _ocr_tesseract(pix, *, lang: str = "vie+eng") -> str:
    import pytesseract
    from PIL import Image

    img = _preprocess(Image.open(io.BytesIO(pix.tobytes("png"))))
    raw = pytesseract.image_to_string(
        img,
        lang=lang,
        config="--oem 3 --psm 6 -c preserve_interword_spaces=1",
    )
    return fix_vietnamese_ocr(str(raw)).strip()


def _ocr_rapid(pix) -> str:
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR

    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
        pix.height, pix.width, pix.n
    )
    if pix.n == 4:
        arr = arr[:, :, :3]
    result, _ = RapidOCR()(arr)
    if not result:
        return ""
    raw = "\n".join(line[1] for line in result if line[1].strip())
    return fix_vietnamese_ocr(raw)


def extract_pdf_text_layer(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"--- Trang {i} ---\n{text}")
    return "\n\n".join(pages)


def ocr_pdf(
    path: Path,
    *,
    max_pages: int | None = None,
    scale: float = 2.5,
    progress: ProgressFn | None = None,
) -> str:
    import pymupdf

    use_tess = configure_tesseract()
    doc = pymupdf.open(str(path))
    limit = min(len(doc), max_pages) if max_pages else len(doc)
    pages: list[str] = []
    try:
        for i in range(limit):
            if progress:
                progress(i + 1, limit)
            page = doc[i]
            pix = page.get_pixmap(
                matrix=pymupdf.Matrix(scale, scale),
                clip=_page_clip(page),
            )
            text = _ocr_tesseract(pix) if use_tess else _ocr_rapid(pix)
            if text.strip():
                pages.append(f"--- Trang {i + 1} ---\n{text.strip()}")
    finally:
        doc.close()
    return "\n\n".join(pages)


def read_pdf(
    file_path: str | Path,
    *,
    use_ocr: bool = True,
    max_pages: int | None = None,
    ocr_scale: float = 2.5,
    progress: ProgressFn | None = None,
) -> str:
    path = Path(file_path)
    text = extract_pdf_text_layer(path)
    if text.strip():
        return text
    if not use_ocr:
        return ""
    return ocr_pdf(
        path, max_pages=max_pages, scale=ocr_scale, progress=progress
    )


def read_word(file_path: str | Path) -> str:
    from docx import Document

    doc = Document(str(file_path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_excel(file_path: str | Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in cast(Any, sheet).iter_rows(values_only=True):
            cells = [
                str(c).strip() for c in row if c is not None and str(c).strip()
            ]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def read_attachment(
    file_path: str | Path,
    *,
    use_ocr: bool = True,
    max_pages: int | None = None,
    progress: ProgressFn | None = None,
) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy file: {path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(
            path, use_ocr=use_ocr, max_pages=max_pages, progress=progress
        )
    if ext == ".docx":
        return read_word(path)
    if ext in {".xlsx", ".xls"}:
        return read_excel(path)
    raise ValueError(
        f"Định dạng không hỗ trợ: {ext}. Chỉ nhận: {', '.join(sorted(SUPPORTED_EXTS))}"
    )
